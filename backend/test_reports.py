"""
Report Generation Tests
=======================
Monthly financial report: income/expense/budget/savings collection,
JSON/CSV/Excel/PDF output, report notifications, empty-data cases,
calculation verification against known records and user isolation.

Run:  cd backend && python -m pytest test_reports.py -v
"""

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.main import app
from app import models

SQLALCHEMY_DATABASE_URL = "sqlite:///./test_reports.db"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False})
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(autouse=True)
def setup_db():
    previous = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    if previous is not None:
        app.dependency_overrides[get_db] = previous
    else:
        app.dependency_overrides.pop(get_db, None)


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def db():
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()


def auth_header(token):
    return {"Authorization": f"Bearer {token}"}


def register(client, username):
    r = client.post("/api/auth/register/", json={
        "username": username,
        "email": f"{username}@test.com",
        "password": "SecurePass123!",
        "password_confirm": "SecurePass123!",
    })
    assert r.status_code == 201
    return auth_header(r.json()["tokens"]["access"])


def add_income(client, headers, amount, source, date):
    r = client.post("/api/incomes/", json={
        "amount": amount, "source": source, "date": date,
    }, headers=headers)
    assert r.status_code == 201


def add_expense(client, headers, amount, category, date):
    r = client.post("/api/expenses/", json={
        "amount": amount, "category": category, "date": date,
    }, headers=headers)
    assert r.status_code == 201


def get_report(client, headers, month=9, year=2026, format="json"):
    return client.get(
        "/api/reports/monthly/",
        headers=headers,
        params={"month": month, "year": year, "format": format},
    )


def seed_september(client, headers):
    """Known September 2026 records for manual verification."""
    add_income(client, headers, 10000, "scholarship", "2026-09-01")
    add_income(client, headers, 5000, "freelance", "2026-09-05")
    add_income(client, headers, 2000, "pocket_money", "2026-08-15")   # outside month
    add_expense(client, headers, 1500, "food", "2026-09-02")
    add_expense(client, headers, 800, "food", "2026-09-06")
    add_expense(client, headers, 1200, "travel", "2026-09-12")
    add_expense(client, headers, 500, "shopping", "2026-08-20")       # outside month

    client.post("/api/budgets/", json={
        "month": 9, "year": 2026, "total_amount": 10000,
        "allocations": [
            {"category": "food", "amount": 5000},
            {"category": "travel", "amount": 2000},
        ],
    }, headers=headers)

    r = client.post("/api/savings/", json={
        "name": "Emergency fund", "target_amount": 50000,
    }, headers=headers)
    goal_id = r.json()["id"]
    client.put(f"/api/savings/{goal_id}/progress/", json={
        "current_saved": 10000,
    }, headers=headers)


# ─── Manual Calculation Verification ─────────────────────────

class TestKnownRecordCalculations:

    def test_totals_match_known_records(self, client):
        headers = register(client, "rep_known")
        seed_september(client, headers)

        r = get_report(client, headers, month=9, year=2026)
        assert r.status_code == 200
        d = r.json()

        # September income: 10000 + 5000 = 15000 (August excluded)
        assert d["income"]["total"] == 15000.0
        assert d["income"]["count"] == 2
        # September expenses: 1500 + 800 + 1200 = 3500 (August excluded)
        assert d["expenses"]["total"] == 3500.0
        assert d["expenses"]["count"] == 3
        # Net: 15000 - 3500 = 11500; rate: 11500/15000*100 = 76.67
        assert d["net_savings"] == 11500.0
        assert d["savings_rate_percent"] == 76.67
        assert d["transaction_count"] == 5
        assert d["month"] == 9
        assert d["year"] == 2026
        assert d["period_start"] == "2026-09-01"
        assert d["period_end"] == "2026-09-30"

    def test_income_breakdown_known_values(self, client):
        headers = register(client, "rep_inc")
        seed_september(client, headers)
        d = get_report(client, headers, month=9, year=2026).json()
        breakdown = {b["label"]: b for b in d["income"]["breakdown"]}

        # scholarship 10000/15000 = 66.67%, freelance 5000/15000 = 33.33%
        assert breakdown["scholarship"]["total"] == 10000.0
        assert breakdown["scholarship"]["count"] == 1
        assert breakdown["scholarship"]["percent"] == 66.67
        assert breakdown["freelance"]["total"] == 5000.0
        assert breakdown["freelance"]["percent"] == 33.33
        assert len(d["income"]["breakdown"]) == 2

    def test_expense_breakdown_known_values(self, client):
        headers = register(client, "rep_exp")
        seed_september(client, headers)
        d = get_report(client, headers, month=9, year=2026).json()
        breakdown = {b["label"]: b for b in d["expenses"]["breakdown"]}

        # food 1500+800=2300 (65.71%), travel 1200 (34.29%)
        assert breakdown["food"]["total"] == 2300.0
        assert breakdown["food"]["count"] == 2
        assert breakdown["food"]["percent"] == 65.71
        assert breakdown["travel"]["total"] == 1200.0
        assert breakdown["travel"]["percent"] == 34.29
        assert len(d["expenses"]["breakdown"]) == 2

    def test_budget_section_known_values(self, client):
        headers = register(client, "rep_bud")
        seed_september(client, headers)
        d = get_report(client, headers, month=9, year=2026).json()

        b = d["budget"]
        assert b is not None
        assert b["total_budget"] == 10000.0
        assert b["total_spent"] == 3500.0
        assert b["remaining"] == 6500.0
        assert b["utilization_percent"] == 35.0

        allocs = {a["category"]: a for a in b["allocations"]}
        # food: 2300/5000 = 46%
        assert allocs["food"]["budgeted"] == 5000.0
        assert allocs["food"]["spent"] == 2300.0
        assert allocs["food"]["remaining"] == 2700.0
        assert allocs["food"]["utilization_percent"] == 46.0
        # travel: 1200/2000 = 60%
        assert allocs["travel"]["spent"] == 1200.0
        assert allocs["travel"]["utilization_percent"] == 60.0

    def test_savings_section_known_values(self, client):
        headers = register(client, "rep_sav")
        seed_september(client, headers)
        d = get_report(client, headers, month=9, year=2026).json()

        s = d["savings"]
        assert s["total_goals"] == 1
        assert s["total_target"] == 50000.0
        assert s["total_saved"] == 10000.0
        assert s["overall_progress_percent"] == 20.0
        assert len(s["goals"]) == 1
        assert s["goals"][0]["name"] == "Emergency fund"
        assert s["goals"][0]["progress_percent"] == 20.0

    def test_formula_consistency(self, client):
        headers = register(client, "rep_formula")
        seed_september(client, headers)
        d = get_report(client, headers, month=9, year=2026).json()
        assert d["net_savings"] == round(
            d["income"]["total"] - d["expenses"]["total"], 2
        )
        assert d["transaction_count"] == d["income"]["count"] + d["expenses"]["count"]

    def test_report_matches_database_records(self, client, db):
        headers = register(client, "rep_db")
        seed_september(client, headers)

        user = db.query(models.User).filter(
            models.User.username == "rep_db",
        ).first()
        db_income = [
            i for i in db.query(models.Income).filter(models.Income.user_id == user.id).all()
            if i.date.year == 2026 and i.date.month == 9
        ]
        db_expenses = [
            e for e in db.query(models.Expense).filter(models.Expense.user_id == user.id).all()
            if e.date.year == 2026 and e.date.month == 9
        ]

        d = get_report(client, headers, month=9, year=2026).json()
        assert d["income"]["total"] == sum(float(i.amount) for i in db_income)
        assert d["expenses"]["total"] == sum(float(e.amount) for e in db_expenses)
        assert d["income"]["count"] == len(db_income)
        assert d["expenses"]["count"] == len(db_expenses)

    def test_other_month_report_excludes_september(self, client):
        headers = register(client, "rep_other")
        seed_september(client, headers)

        d = get_report(client, headers, month=8, year=2026).json()
        assert d["income"]["total"] == 2000.0
        assert d["expenses"]["total"] == 500.0
        assert d["net_savings"] == 1500.0
        assert d["budget"] is None  # no August budget


# ─── Empty Data ──────────────────────────────────────────────

class TestEmptyData:

    def test_empty_month_returns_zeros(self, client):
        headers = register(client, "rep_empty")
        d = get_report(client, headers, month=1, year=2026).json()

        assert d["income"]["total"] == 0.0
        assert d["income"]["count"] == 0
        assert d["income"]["breakdown"] == []
        assert d["expenses"]["total"] == 0.0
        assert d["expenses"]["count"] == 0
        assert d["expenses"]["breakdown"] == []
        assert d["net_savings"] == 0.0
        assert d["savings_rate_percent"] == 0.0
        assert d["budget"] is None
        assert d["transaction_count"] == 0
        assert d["savings"]["total_goals"] == 0

    def test_empty_month_csv(self, client):
        headers = register(client, "rep_empty_csv")
        r = get_report(client, headers, month=1, year=2026, format="csv")
        assert r.status_code == 200
        assert "text/csv" in r.headers["content-type"]
        assert "BudgetBuddy Monthly Report" in r.text
        assert "No budget set" in r.text

    def test_empty_month_pdf(self, client):
        headers = register(client, "rep_empty_pdf")
        r = get_report(client, headers, month=1, year=2026, format="pdf")
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"

    def test_expense_only_month(self, client):
        headers = register(client, "rep_expense_only")
        add_expense(client, headers, 750, "food", "2026-03-10")

        d = get_report(client, headers, month=3, year=2026).json()
        assert d["income"]["total"] == 0.0
        assert d["expenses"]["total"] == 750.0
        assert d["net_savings"] == -750.0
        assert d["savings_rate_percent"] == 0.0


# ─── Output Formats ──────────────────────────────────────────

class TestOutputFormats:

    def test_json_format_default(self, client):
        headers = register(client, "rep_fmt_json")
        seed_september(client, headers)
        r = get_report(client, headers)
        assert r.status_code == 200
        assert "application/json" in r.headers["content-type"]
        d = r.json()
        assert d["month"] == 9

    def test_csv_format(self, client):
        headers = register(client, "rep_fmt_csv")
        seed_september(client, headers)
        r = get_report(client, headers, format="csv")
        assert r.status_code == 200
        assert "text/csv" in r.headers["content-type"]
        assert "attachment" in r.headers.get("content-disposition", "")
        assert "budgetbuddy_report_2026-09.csv" in r.headers["content-disposition"]
        text = r.text
        assert "BudgetBuddy Monthly Report" in text
        assert "Income Total,15000.00" in text
        assert "Expense Total,3500.00" in text
        assert "Net Savings,11500.00" in text
        assert "scholarship,10000.00" in text
        assert "food,2300.00" in text
        assert "Total Budget,10000.00" in text

    def test_xlsx_format(self, client):
        headers = register(client, "rep_fmt_xlsx")
        seed_september(client, headers)
        r = get_report(client, headers, format="xlsx")
        assert r.status_code == 200
        assert (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            in r.headers["content-type"]
        )
        assert "budgetbuddy_report_2026-09.xlsx" in r.headers["content-disposition"]
        # XLSX is a ZIP container
        assert r.content[:2] == b"PK"

        # Verify it opens with openpyxl
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        rows = [row for row in ws.iter_rows(values_only=True)]
        flat = [str(c) for row in rows for c in row if c is not None]
        assert any("BudgetBuddy Monthly Report" in c for c in flat)
        assert any(c == 15000 for c in rows[1] if c is not None) or any(
            "15000" in str(c) for c in flat
        )

    def test_pdf_format(self, client):
        headers = register(client, "rep_fmt_pdf")
        seed_september(client, headers)
        r = get_report(client, headers, format="pdf")
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert "attachment" in r.headers.get("content-disposition", "")
        assert "budgetbuddy_report_2026-09.pdf" in r.headers["content-disposition"]
        assert r.content[:5] == b"%PDF-"
        assert len(r.content) > 500

    def test_unsupported_format_rejected(self, client):
        headers = register(client, "rep_fmt_bad")
        r = get_report(client, headers, format="docx")
        assert r.status_code == 400
        assert "Unsupported format" in r.json()["detail"]

    def test_format_excel_alias_maps_to_xlsx(self, client):
        headers = register(client, "rep_fmt_excel")
        seed_september(client, headers)
        r = get_report(client, headers, format="excel")
        assert r.status_code == 200
        assert (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            in r.headers["content-type"]
        )
        assert "budgetbuddy_report_2026-09.xlsx" in r.headers["content-disposition"]
        assert r.content[:2] == b"PK"

    def test_format_xls_alias_maps_to_xlsx(self, client):
        headers = register(client, "rep_fmt_xls")
        r = get_report(client, headers, format="xls")
        assert r.status_code == 200
        assert r.content[:2] == b"PK"

    def test_format_is_case_insensitive(self, client):
        headers = register(client, "rep_fmt_case")
        r = get_report(client, headers, format="CSV")
        assert r.status_code == 200
        assert "text/csv" in r.headers["content-type"]

    def test_format_pdf_uppercase(self, client):
        headers = register(client, "rep_fmt_pdf_up")
        r = get_report(client, headers, format="PDF")
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"


# ─── Invalid Inputs ──────────────────────────────────────────

class TestInvalidInput:

    def test_invalid_month(self, client):
        headers = register(client, "rep_inv1")
        r = get_report(client, headers, month=13, year=2026)
        assert r.status_code == 400
        assert "month" in r.json()["detail"]

    def test_month_zero(self, client):
        headers = register(client, "rep_inv2")
        r = get_report(client, headers, month=0, year=2026)
        assert r.status_code == 400

    def test_negative_month(self, client):
        headers = register(client, "rep_inv3")
        r = get_report(client, headers, month=-5, year=2026)
        assert r.status_code == 400

    def test_invalid_year(self, client):
        headers = register(client, "rep_inv4")
        r = get_report(client, headers, month=9, year=0)
        assert r.status_code == 400

    def test_missing_month(self, client):
        headers = register(client, "rep_inv5")
        r = client.get("/api/reports/monthly/?year=2026", headers=headers)
        assert r.status_code == 422

    def test_missing_year(self, client):
        headers = register(client, "rep_inv6")
        r = client.get("/api/reports/monthly/?month=9", headers=headers)
        assert r.status_code == 422

    def test_non_integer_month(self, client):
        headers = register(client, "rep_inv7")
        r = client.get(
            "/api/reports/monthly/",
            headers=headers,
            params={"month": "abc", "year": 2026},
        )
        assert r.status_code == 422


# ─── Report Notifications ────────────────────────────────────

class TestReportNotifications:

    def test_report_generation_creates_notification(self, client):
        headers = register(client, "rep_notif")
        seed_september(client, headers)

        r = get_report(client, headers, month=9, year=2026)
        assert r.status_code == 200

        notifications = client.get("/api/notifications/", headers=headers).json()["notifications"]
        report_notifs = [
            n for n in notifications
            if "Monthly report" in n["message"]
        ]
        assert len(report_notifs) == 1
        assert "09/2026" in report_notifs[0]["message"]
        assert report_notifs[0]["notification_type"] == "general"

    def test_report_notification_not_duplicated(self, client):
        headers = register(client, "rep_notif_dup")
        seed_september(client, headers)

        get_report(client, headers, month=9, year=2026)
        get_report(client, headers, month=9, year=2026)
        get_report(client, headers, month=9, year=2026, format="csv")

        notifications = client.get("/api/notifications/", headers=headers).json()["notifications"]
        report_notifs = [
            n for n in notifications
            if "Monthly report" in n["message"] and "09/2026" in n["message"]
        ]
        assert len(report_notifs) == 1

    def test_separate_notifications_per_month(self, client):
        headers = register(client, "rep_notif_months")
        seed_september(client, headers)
        add_expense(client, headers, 100, "food", "2026-10-05")

        get_report(client, headers, month=9, year=2026)
        get_report(client, headers, month=10, year=2026)

        notifications = client.get("/api/notifications/", headers=headers).json()["notifications"]
        report_notifs = [n for n in notifications if "Monthly report" in n["message"]]
        assert len(report_notifs) == 2

    def test_notification_respects_profile_preference(self, client, db):
        headers = register(client, "rep_notif_off")
        seed_september(client, headers)

        user = db.query(models.User).filter(
            models.User.username == "rep_notif_off",
        ).first()
        profile = db.query(models.Profile).filter(
            models.Profile.user_id == user.id,
        ).first()
        profile.monthly_report_notifications = False
        db.commit()

        get_report(client, headers, month=9, year=2026)

        notifications = client.get("/api/notifications/", headers=headers).json()["notifications"]
        report_notifs = [n for n in notifications if "Monthly report" in n["message"]]
        assert report_notifs == []


# ─── Auth & User Isolation ───────────────────────────────────

class TestAuthAndIsolation:

    def test_unauthenticated_rejected(self, client):
        r = client.get("/api/reports/monthly/?month=9&year=2026")
        assert r.status_code in (401, 403)

    def test_invalid_token_rejected(self, client):
        r = client.get(
            "/api/reports/monthly/?month=9&year=2026",
            headers={"Authorization": "Bearer bad.token"},
        )
        assert r.status_code in (401, 403)

    def test_user_sees_only_own_report_data(self, client):
        ha = register(client, "rep_iso_a")
        hb = register(client, "rep_iso_b")

        add_income(client, ha, 10000, "scholarship", "2026-09-01")
        add_expense(client, ha, 3000, "food", "2026-09-10")
        add_income(client, hb, 500, "freelance", "2026-09-03")

        da = get_report(client, ha, month=9, year=2026).json()
        db_ = get_report(client, hb, month=9, year=2026).json()

        assert da["income"]["total"] == 10000.0
        assert da["expenses"]["total"] == 3000.0
        assert db_["income"]["total"] == 500.0
        assert db_["expenses"]["total"] == 0.0

    def test_notifications_isolated(self, client):
        ha = register(client, "rep_notif_iso_a")
        hb = register(client, "rep_notif_iso_b")
        seed_september(client, ha)

        get_report(client, ha, month=9, year=2026)

        notifs_a = client.get("/api/notifications/", headers=ha).json()["notifications"]
        notifs_b = client.get("/api/notifications/", headers=hb).json()["notifications"]

        assert any("Monthly report" in n["message"] for n in notifs_a)
        assert not any("Monthly report" in n["message"] for n in notifs_b)

    def test_report_structure_keys(self, client):
        headers = register(client, "rep_shape")
        d = get_report(client, headers, month=9, year=2026).json()
        expected = {
            "month", "year", "generated_at", "period_start", "period_end",
            "income", "expenses", "net_savings", "savings_rate_percent",
            "budget", "savings", "transaction_count",
            "income_count", "expense_count",
        }
        assert expected == set(d.keys())
        assert set(d["income"].keys()) == {"total", "count", "breakdown"}
        assert set(d["expenses"].keys()) == {"total", "count", "breakdown"}
        assert set(d["savings"].keys()) == {
            "total_goals", "total_target", "total_saved",
            "overall_progress_percent", "goals",
        }


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v", "--tb=short"]))
